"""Write the review workbook a landman works from (pilot stand-in for the review screen)."""
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

HEAD = PatternFill("solid", fgColor="1F3A5F")
SEVERITY_FILL = {"high": "F8D7DA", "medium": "FFF2CC", "low": "E2EFDA", "review": "DDEBF7"}


def _table(ws, headers, rows, widths=None):
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEAD
    for r in rows:
        ws.append([", ".join(v) if isinstance(v, list) else v for v in r])
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def export_review(path, intake, run_id, chain, flags, skipped, problems) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "Summary"
    for k, v in [("Run", run_id), ("Tract", intake.tract_id), ("County", intake.county),
                 ("Legal description", intake.legal_description), ("Title agent", intake.agent),
                 ("Package source", intake.mode + (f" ({intake.runsheet_path.name})" if intake.runsheet_path else "")),
                 ("Tract acreage", f"{chain['subject_acres']:g}" if chain.get("subject_acres") else "unknown"),
                 ("Current surface owner", ", ".join(chain["current_owner"]) or "NOT FOUND"),
                 ("Ownership notes", len(chain.get("ownership_notes", [])) or "none"),
                 ("Vesting instrument", chain["vesting_instrument"]),
                 ("Open encumbrances", sum(e["status"] == "open" for e in chain["encumbrances"])),
                 ("Flags", len(flags)),
                 ("Instruments with an unrecognized type", ", ".join(chain["unclassified"]) or "none"),
                 ("Instruments with no filing date", ", ".join(chain["undated"]) or "none"),
                 ("Supporting documents (not in the chain)", len(chain.get("supporting", []))),
                 ("Rules not yet implemented (skipped)", ", ".join(skipped) or "none"),
                 ("Package problems", "; ".join(problems) or "none")]:
        ws.append([k, v])
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 80

    _table(wb.create_sheet("Chain"),
           ["Filing date", "Type", "Instrument", "Grantors", "Grantees", "Acres", "Interest", "File"],
           [[i["filing_date"], i["type"], i["instrument_id"], i["grantors"], i["grantees"], i.get("acreage"),
             i.get("interest"), i["file_name"]] for i in chain["instruments"]], [12, 24, 22, 32, 32, 8, 28, 30])
    _table(wb.create_sheet("Encumbrances"), ["Filing date", "Type", "Instrument", "Holder", "Status"],
           [[e["filing_date"], e["type"], e["instrument_id"], e["holder"], e["status"]]
            for e in chain["encumbrances"]], [12, 24, 22, 32, 12])

    _table(wb.create_sheet("Supporting docs"), ["Kind", "File"],
           [[i.get("kind"), i["file_name"]] for i in chain.get("supporting", [])], [30, 60])

    fs = wb.create_sheet("Flags")
    _table(fs, ["Rule", "Severity", "Instrument", "Message", "Reviewer action", "Note"],
           [[f.rule_id, f.severity, f.instrument_id, f.message, None, None] for f in flags],
           [8, 12, 22, 70, 16, 40])
    for row in fs.iter_rows(min_row=2):
        color = SEVERITY_FILL.get(str(row[1].value).split()[0].lower())
        if color:
            row[1].fill = PatternFill("solid", fgColor=color)
    wb.save(path)
    return path
