"""Read the title agent's intake workbook and check the package is complete."""
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path

import openpyxl
from .fuzzy import token_sort_ratio

DOC_EXTS = {".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff"}

# Intake Form labels (column A) -> attribute. Matched by "starts with".
LABELS = {
    "tract id": "tract_id",
    "project / client": "project",
    "county": "county",
    "legal description": "legal_description",
    "search period": "search_period",
    "owners searched by name": "owners",
    "records searched": "records_searched",
    "title agent": "agent",
    "delivery date": "delivery_date",
    "agent notes": "notes",
}


def _clean(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    s = str(v).strip()
    if not s or s.lower().startswith("example"):
        return None
    return s


@dataclass
class DocRow:
    file_name: str
    instrument_type: str | None = None
    recording_ref: str | None = None
    filing_date: str | None = None
    grantor: str | None = None
    grantee: str | None = None
    found_by: str | None = None
    owner_searched: str | None = None
    notes: str | None = None


@dataclass
class Intake:
    tract_id: str | None = None
    project: str | None = None
    county: str | None = None
    legal_description: str | None = None
    search_period: str | None = None
    owners: list = field(default_factory=list)
    records_searched: str | None = None
    agent: str | None = None
    delivery_date: str | None = None
    notes: str | None = None
    documents: list = field(default_factory=list)
    source: Path | None = None


def find_intake_file(folder) -> Path | None:
    for p in sorted(Path(folder).glob("*.xlsx")):
        if p.name.startswith("~$"):
            continue
        try:
            wb = openpyxl.load_workbook(p, read_only=True)
        except Exception:
            continue
        if "Intake Form" in wb.sheetnames and "Document List" in wb.sheetnames:
            return p
    return None


def read_intake(path) -> Intake:
    wb = openpyxl.load_workbook(path, data_only=True)
    intake = Intake(source=Path(path))

    for row in wb["Intake Form"].iter_rows(values_only=True):
        if not row or row[0] is None or len(row) < 2:
            continue
        label = str(row[0]).strip().lower()
        for key, attr in LABELS.items():
            if label.startswith(key):
                value = _clean(row[1])
                if attr == "owners":
                    owners = [re.sub(r"\(.*?\)", "", o).strip() for o in (value or "").split(";")]
                    intake.owners = [o for o in owners if o]
                else:
                    setattr(intake, attr, value)
                break

    ws = wb["Document List"]
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not _clean(row[0]):
            continue
        if row[8] and str(row[8]).strip().lower().startswith("example"):
            continue  # template example rows
        intake.documents.append(DocRow(*[_clean(v) for v in row[:9]]))
    return intake


def check_package(folder, intake: Intake) -> list[str]:
    """Problems that should go back to the title agent before extraction."""
    folder = Path(folder)
    problems = []
    if not intake.tract_id:
        problems.append("Intake form has no tract ID")
    elif folder.name != intake.tract_id:
        problems.append(f"Folder name '{folder.name}' does not match tract ID '{intake.tract_id}'")
    if not intake.documents:
        problems.append("Document List is empty")

    files = {p.name for p in folder.iterdir() if p.suffix.lower() in DOC_EXTS}
    listed = {d.file_name for d in intake.documents}
    problems += [f"Listed but not uploaded: {f}" for f in sorted(listed - files)]
    problems += [f"Uploaded but not on the Document List: {f}" for f in sorted(files - listed)]

    if not intake.owners:
        problems.append("No owners listed under 'Owners searched by name'")
    for d in intake.documents:
        if (d.found_by or "").lower() != "name search":
            continue
        if not d.owner_searched:
            problems.append(f"{d.file_name}: name search with no owner name")
        elif not any(token_sort_ratio(d.owner_searched.lower(), o.lower()) >= 85 for o in intake.owners):
            problems.append(f"{d.file_name}: owner '{d.owner_searched}' is not in the intake owners list")
    return problems
