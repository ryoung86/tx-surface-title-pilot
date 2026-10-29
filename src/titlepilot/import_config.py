"""Turn the landmen's Fields and Rules workbook into config/fields.yaml and rules.yaml.

Only rows with Status 'Starter' or 'Approved' are imported; 'Proposed' and
'Rejected' rows stay in the workbook until the pilot lead decides.
Settings that live only in rules.yaml (enabled, params) are kept across imports.
"""
from pathlib import Path

import openpyxl
import yaml

USE_STATUS = {"starter", "approved"}

# Rules implemented in code that are not in the landmen's workbook.
CODE_RULES = [
    {
        "id": "PK01",
        "name": "Owner in the chain has no record of a name search (intake form or runsheet)",
        "severity": "High",
        "enabled": True,
        "source": "code",
    },
    {
        "id": "PK02",
        "name": "Instrument on the agent's runsheet was not uploaded",
        "severity": "High",
        "enabled": True,
        "source": "code",
    },
    {
        "id": "PK03",
        "name": "Uploaded document is not on the agent's runsheet",
        "severity": "Low",
        "enabled": True,
        "source": "code",
    },
]


def _clean(v):
    if v is None:
        return None
    s = str(v).strip()
    if not s or s.lower().startswith("example"):
        return None
    return s


def _rows(ws):
    headers = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        yield dict(zip(headers, row))


def import_workbook(xlsx, config_dir="config") -> dict:
    wb = openpyxl.load_workbook(xlsx, data_only=True)
    out = Path(config_dir)
    out.mkdir(parents=True, exist_ok=True)

    fields = []
    for r in _rows(wb["Extraction Fields"]):
        name = _clean(r.get("Field"))
        if not name or (r.get("Status") or "").strip().lower() not in USE_STATUS:
            continue
        for n in (x.strip() for x in name.split(",") if x.strip()):
            fields.append({
                "name": n,
                "group": (_clean(r.get("Group")) or "").lower(),
                "type": _clean(r.get("Type")) or "text",
                "description": _clean(r.get("Definition")) or "",
                "required": (str(r.get("Required?") or "")).strip().lower() == "yes",
            })

    rules_path = out / "rules.yaml"
    existing = {}
    if rules_path.exists():
        existing = {x["id"]: x for x in (yaml.safe_load(rules_path.read_text()) or {}).get("rules", [])}

    rules = []
    for r in _rows(wb["Rules"]):
        rid = _clean(r.get("ID"))
        if not rid or (r.get("Status") or "").strip().lower() not in USE_STATUS:
            continue
        old = existing.get(rid, {})
        entry = {
            "id": rid,
            "name": _clean(r.get("Rule")) or "",
            "severity": _clean(r.get("Severity")) or "Review",
            "enabled": old.get("enabled", True),
        }
        applies = _clean(r.get("Applies to (instrument types)"))
        if applies:
            entry["applies_to"] = [x.strip().lower() for x in applies.split(",") if x.strip()]
        ref = _clean(r.get("Legal reference"))
        if ref:
            entry["legal_reference"] = ref
        if old.get("params"):
            entry["params"] = old["params"]
        rules.append(entry)

    ids = {x["id"] for x in rules}
    for cr in CODE_RULES:
        if cr["id"] not in ids:
            rules.append(existing.get(cr["id"], cr))

    header = "# Generated from {} by `titlepilot import-config`.\n# Edit the workbook, not this file, except for 'enabled' and 'params'.\n".format(Path(xlsx).name)
    (out / "fields.yaml").write_text(header + yaml.safe_dump({"fields": fields}, sort_keys=False, allow_unicode=True, width=1000))
    rules_path.write_text(header + yaml.safe_dump({"rules": rules}, sort_keys=False, allow_unicode=True, width=1000))
    return {"fields": len(fields), "rules": len(rules)}
