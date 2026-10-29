"""Write the title report in SEG's format, from the blank template in templates/.

Sheets filled, matching the SEG report layout:
  Ownership          header block, property description, surface and pore space ownership
  Ownership BT Data  one row per owner
  Encumbrances       mortgages/liens table, roads, railroads, all other current encumbrances
  Flowchart          the chain of title drawn in cells (the template's drawn boxes can't be copied)
  Names Examined     owners in the chain with their ownership periods
  Index to All Documents, Runsheet   every recorded instrument in filing order
Mineral, NPRI and non-executive sections are left blank: this is a surface report.
Three AI sheets are added at the end: AI Review (summary and ownership notes), AI Flags,
AI Supporting Docs. Everything here is a draft for a landman to check and sign.
"""
import re
from copy import copy
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from .validate import norm_ref

TEMPLATE = Path(__file__).with_name("templates") / "surface_title_report.xlsx"

ABBREVIATIONS = [  # (keyword in the instrument type, abbreviation used on SEG runsheets); first match wins
    ("special warranty", "SWD"), ("warranty deed", "WD"), ("deed of gift", "GD"), ("gift deed", "GD"),
    ("quitclaim", "QCD"), ("deed without warranty", "DWW"), ("deed of trust", "DT"), ("release", "Rel"),
    ("right-of-way", "ROW"), ("right of way", "ROW"), ("easement", "Esmt"), ("ucc", "UCC"),
    ("financing statement", "UCC"), ("abstract of judgment", "AJ"), ("judgment", "Jdgmt"), ("lis pendens", "LP"),
    ("affidavit of heirship", "AOH"), ("correction", "Corr. Affd."), ("affidavit", "Affd."),
    ("oil, gas", "OGML"), ("oil and gas", "OGML"), ("mineral lease", "OGML"), ("probate", "Prob."),
    ("tax deed", "Tax Deed"), ("sheriff", "Sheriff's Deed"), ("lien", "Lien"), ("plat", "Plat"),
]

RED = PatternFill("solid", fgColor="F8D7DA")


def abbrev(itype: str | None) -> str:
    t = (itype or "").lower()
    for key, short in ABBREVIATIONS:
        if key in t:
            return short
    return (itype or "").title()


def display_ref(inst) -> str:
    """'2011-9176' -> '2011009176', 'V708P1' -> '708/1', as written on SEG runsheets."""
    key = norm_ref(inst.get("instrument_id"))
    m = re.fullmatch(r"(\d{4})-(\d+)", key)
    if m:
        return f"{m.group(1)}{int(m.group(2)):06d}"
    m = re.fullmatch(r"V(\d+)P(\d+)", key)
    if m:
        return f"{m.group(1)}/{m.group(2)}"
    return inst.get("instrument_id") or ""


def record_type(inst) -> str:
    """Deed Records for volume/page references, Official Public Records for instrument numbers."""
    return "DR" if norm_ref(inst.get("instrument_id")).startswith("V") else "OPR"


def as_date(iso):
    try:
        return date.fromisoformat(iso) if iso else None
    except ValueError:
        return None


def fmt_date(iso) -> str:
    d = as_date(iso)
    return f"{d.month}/{d.day}/{d.year}" if d else "undated"


def names(lst) -> str:
    return ", ".join(n for n in (lst or []) if n)


def abstractor_note(inst) -> str:
    """Short note in the style of the SEG runsheet: what it conveys, what it reserves."""
    parts = []
    cat = inst.get("category")
    if cat == "conveyance":
        parts.append(f"Conveys {inst['interest']}" if inst.get("interest") else "Conveys")
        if inst.get("mineral_exception"):
            parts.append("reserves/excepts minerals")
    elif cat == "easement":
        parts.append(inst.get("easements") or "Easement")
    elif cat in ("lien", "judgment"):
        parts.append(inst.get("liens") or inst.get("type", "").title())
    elif inst.get("notes"):
        parts.append(inst["notes"])
    return "; ".join(p for p in parts if p)[:250]


def _copy_row_style(ws, src_row, dst_row, c1, c2):
    for c in range(c1, c2 + 1):
        ws.cell(dst_row, c)._style = copy(ws.cell(src_row, c)._style)
    ws.row_dimensions[dst_row].height = ws.row_dimensions[src_row].height


def _fit_height(ws, row, c1, c2, min_height=24, pt=9):
    """Row height from the longest wrapped value in the row (template heights belong to other rows)."""
    lines = 1
    for c in range(c1, c2 + 1):
        v = ws.cell(row, c).value
        if v is None or not isinstance(v, str):
            continue
        width = ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width or 10
        per_line = max(8, int(width * 1.15))
        lines = max(lines, sum(1 + len(part) // per_line for part in v.split("\n")))
    ws.row_dimensions[row].height = max(min_height, lines * (pt + 3) + 4)


def _fit_width(ws, landscape=None):
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    if landscape is not None:
        ws.page_setup.orientation = "landscape" if landscape else "portrait"


def _write_rows(ws, first_row, c1, c2, rows):
    """Write rows starting at first_row, styling any rows beyond the template's like the first one."""
    styled_until = first_row
    while styled_until <= ws.max_row and ws.cell(styled_until, c1).has_style:
        styled_until += 1
    for n, values in enumerate(rows):
        r = first_row + n
        if r >= styled_until:
            _copy_row_style(ws, first_row, r, c1, c2)
        for c, v in enumerate(values, c1):
            ws.cell(r, c).value = v
        _fit_height(ws, r, c1, c2)
    return first_row + len(rows) - 1 if rows else first_row


def _owner_periods(chain):
    """When each owner in the chain took and parted with title (for Names Examined).
    Out-sales and undivided-interest deeds don't end a period; current owners run to the end."""
    from .chain import is_fractional, is_partial_area
    periods = {}
    for c in chain["conveyances"]:
        for g in c["grantees"]:
            periods.setdefault(g, [c["filing_date"], None])
        if is_fractional(c) or is_partial_area(c, chain.get("subject_acres")):
            continue
        for g in c["grantors"]:
            p = periods.setdefault(g, [None, None])
            p[1] = c["filing_date"]
    for o in chain["current_owner"]:
        if o in periods:
            periods[o][1] = None
    return periods


def _split_name(n):
    if re.search(r"\b(llc|inc|co|company|corp|corporation|trust|ltd|lp|partnership|bank|estate)\b\.?", n, re.I):
        return n, "", ""
    parts = n.replace(",", " ").split()
    suffix = [p for p in parts if p.rstrip(".").lower() in ("jr", "sr", "ii", "iii", "iv")]
    core = [p for p in parts if p not in suffix]
    if len(core) < 2:
        return n, "", ""
    return " ".join(core[-1:] + suffix), core[0], " ".join(core[1:-1])


def write_report(path, intake, chain, flags, skipped, problems, run_id, researched=None) -> Path:
    """researched: (from_iso, to_iso) for the search period, if known."""
    wb = openpyxl.load_workbook(TEMPLATE)
    instruments = chain["instruments"]
    acres = chain.get("subject_acres")
    owners = chain["current_owner"]
    vesting = next((i for i in instruments if i["instrument_id"] == chain.get("vesting_instrument")), None)
    dated = [i["filing_date"] for i in instruments if i["filing_date"]]
    r_from, r_to = researched or (min(dated) if dated else None, max(dated) if dated else None)
    client = (intake.tract_id or "").split(".")[0].upper()
    legal = intake.legal_description or (vesting or {}).get("legal") or ""

    # ---- Ownership ----------------------------------------------------------------------------
    ws = wb["Ownership "]
    if client:
        ws["B5"] = f"{client} TITLE OWNERSHIP REPORT"
    ws["B8"] = f"REPORT DATE: {date.today().month}/{date.today().day}/{date.today().year}"
    ws["B9"] = "REVISED DATE: N/A"
    ws["B11"] = names(owners) or "OWNER NOT DETERMINED"
    ws["C14"] = intake.tract_id
    ws["E14"] = getattr(intake, "aoi", None)
    ws["C15"] = (intake.county or "").replace(" County", "")
    ws["C16"] = intake.agent or "AI draft - landman review required"
    ws["E16"] = acres
    ws["C17"], ws["E17"] = as_date(r_from), as_date(r_to)
    ws["B22"] = legal
    for row in (28, 35):                                   # surface, then pore space (follows the surface)
        ws[f"B{row}"] = "\n".join(owners) if owners else "OWNER NOT DETERMINED"
        ws[f"C{row}"] = display_ref(vesting) if vesting else None
        ws[f"D{row}"] = 1 if owners else 0
        ws[f"E{row}"] = acres
        ws[f"B{row}"].alignment = Alignment(wrap_text=True, vertical="center", horizontal="left")
    ws["B62"] = intake.records_searched or None

    # ---- Ownership BT Data --------------------------------------------------------------------
    ws = wb["Ownership BT Data"]
    for n, o in enumerate(owners):
        last, first, middle = _split_name(o)
        for c, v in enumerate([intake.tract_id, round(100 / len(owners), 6), last, first, middle,
                               None, None, None, None, None, None, "S"], 1):
            ws.cell(2 + n, c).value = v

    # ---- Encumbrances -------------------------------------------------------------------------
    ws = wb["Encumbrances"]
    current = [e for e in chain["encumbrances"] if e["status"] == "open"]
    is_mortgage = lambda e: e["category"] == "lien" and any(k in e["type"] for k in ("deed of trust", "mortgage",
                                                                                     "vendor"))
    text = lambda e: (f"{abbrev(e['type'])} - {record_type(e)} {display_ref(e)}, dated "
                      f"{fmt_date(e.get('execution_date') or e['filing_date'])} from {names(e['grantors']) or '?'} "
                      f"to {names(e['grantees']) or '?'}" + (f" - {abstractor_note(e)}" if abstractor_note(e) else ""))
    mortgages = [e for e in current if is_mortgage(e)]
    roads = [e for e in current if e["category"] == "easement" and "road" in f"{e['type']} {e.get('easements')}".lower()]
    rails = [e for e in current if "rail" in f"{e['type']} {e.get('easements')}".lower() and e not in roads]
    leases = [e for e in current if e["category"] == "surface agreement"]
    other = [e for e in current if e not in mortgages + roads + rails + leases]

    # Burdens created inside conveyances: reserved easements and deed restrictions (rights of first
    # refusal, use covenants). They aren't separate instruments, so they're listed as text lines.
    deed_lines_roads, deed_lines_other = [], []
    for c in chain["conveyances"]:
        tag = f"{abbrev(c['type'])} - {record_type(c)} {display_ref(c)}, dated {fmt_date(c.get('execution_date') or c['filing_date'])}"
        if c.get("easements"):
            line = f"{tag} - reserved: {c['easements']}"
            (deed_lines_roads if "road" in c["easements"].lower() else deed_lines_other).append(line)
        if c.get("restrictions"):
            deed_lines_other.append(f"{tag} - {c['restrictions']}")

    for n, e in enumerate(leases[:9]):
        for c, v in zip("BCDEFGHI", [abbrev(e["type"]), names(e["grantors"]), names(e["grantees"]),
                                     as_date(e.get("execution_date") or e["filing_date"]), None, display_ref(e),
                                     "Open", e.get("legal")]):
            ws[f"{c}{5 + n}"] = v
    ws["C6"] = None if leases else "None"
    for n, e in enumerate(mortgages[:3]):
        for c, v in zip("BCDEFGHI", [abbrev(e["type"]), names(e["grantors"]), names(e["grantees"]),
                                     as_date(e.get("execution_date") or e["filing_date"]), None, display_ref(e), None,
                                     e.get("legal")]):
            ws[f"{c}{18 + n}"] = v
    if not mortgages:
        ws["C18"] = "None"
    other += mortgages[3:] + leases[9:]
    extra = {"B29": deed_lines_roads, "B44": deed_lines_other}
    for cell, items in (("B24", []), ("B29", roads), ("B34", rails), ("B39", []), ("B44", other)):
        lines = [text(e) for e in items] + extra.get(cell, [])
        ws[cell] = "\n".join(lines) if lines else "None"
        ws[cell].alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
        n = max(1, sum(1 + len(t) // 110 for t in lines))
        ws.row_dimensions[ws[cell].row].height = max(ws.row_dimensions[ws[cell].row].height or 15, 15 * n)

    _fit_width(ws)

    # ---- Names Examined -----------------------------------------------------------------------
    ws = wb["Names Examined"]
    periods = _owner_periods(chain)
    searched = [n for n in (intake.owners or [])]
    rows = []
    for name in chain["owners"]:
        start, end = periods.get(name, [None, None])
        rows.append([name, as_date(start), as_date(end) or as_date(r_to)] + [None] * 9)
    for name in searched:
        if not any(name.lower() == r[0].lower() for r in rows):
            rows.append([name, None, None] + [None] * 9)
    _write_rows(ws, 5, 2, 13, rows)

    # ---- Index to All Documents and Runsheet ---------------------------------------------------
    page = 1
    index_rows, run_rows = [], []
    for n, i in enumerate(instruments, 1):
        eff = as_date(i.get("execution_date"))
        filed = as_date(i["filing_date"])
        index_rows.append([page, abbrev(i["type"]), names(i["grantors"]), names(i["grantees"]), eff, filed,
                           display_ref(i) if record_type(i) == "DR" else None, i.get("instrument_no")
                           or (display_ref(i) if record_type(i) == "OPR" else None)])
        page += i.get("pages") or 1
        run_rows.append([n, display_ref(i), abbrev(i["type"]), names(i["grantors"]), names(i["grantees"]), eff,
                         filed, record_type(i), f"{intake.county} County" if intake.county else None,
                         (i.get("legal") or "")[:250], abstractor_note(i)])
    ws = wb["Index to All Documents"]
    last = _write_rows(ws, 3, 2, 9, index_rows)
    ws.conditional_formatting.add(f"F3:G{max(last, 3)}", FormulaRule(formula=["AND($F3<>\"\",$G3<>\"\",$G3<$F3)"], fill=RED))
    ws = wb["Runsheet "]
    a, b = fmt_date(r_from), fmt_date(r_to)
    ws["A4"] = f"Records from {a} to {b}" if researched else f"Documents in package dated {a} to {b}"
    ws["A6"] = legal
    ws["A8"] = f"{acres:.2f} ACRES, MORE OR LESS" if acres else None
    last = _write_rows(ws, 13, 1, 11, run_rows)
    ws.conditional_formatting.add(f"F13:G{max(last, 13)}", FormulaRule(formula=["AND($F13<>\"\",$G13<>\"\",$G13<$F13)"], fill=RED))
    ws.print_area = f"A1:K{max(last, 13)}"

    # ---- Flowchart (chain of title in cells) ---------------------------------------------------
    pos = wb.sheetnames.index("Flowchart")
    wb.remove(wb["Flowchart"])                 # the template's sheet is sized for drawn shapes
    ws = wb.create_sheet("Flowchart", pos)
    thin = Side(style="thin", color="595959")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    ws["B1"] = "CHAIN OF TITLE (surface conveyances, filing order)"
    ws["B1"].font = Font(name="Times New Roman", size=14, bold=True)
    ws.column_dimensions["B"].width = 90
    r = 3
    for c in chain["conveyances"]:
        size = f"{c['acreage']:g} ac" if c.get("acreage") else "acreage not read"
        interest = f"; {c['interest']}" if c.get("interest") else ""
        is_vesting = c["instrument_id"] == chain.get("vesting_instrument")
        cell = ws.cell(r, 2, f"{abbrev(c['type'])} {display_ref(c)}  ({fmt_date(c['filing_date'])})\n"
                             f"{names(c['grantors']) or '?'}  to  {names(c['grantees']) or '?'}\n{size}{interest}"
                             + ("\nVESTING DEED (current owner)" if is_vesting else ""))
        cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        cell.border = box
        cell.font = Font(name="Times New Roman", size=10, bold=is_vesting)
        if is_vesting:
            cell.fill = PatternFill("solid", fgColor="E2EFDA")
        ws.row_dimensions[r].height = 62 if is_vesting else 48
        ws.cell(r + 1, 2, "↓").alignment = Alignment(horizontal="center")
        r += 2
    ws.cell(r, 2, f"CURRENT SURFACE OWNER: {names(owners) or 'not determined'}").font = Font(
        name="Times New Roman", size=12, bold=True)
    ws.print_area = f"B1:B{r}"
    _fit_width(ws, landscape=False)

    # ---- AI sheets ----------------------------------------------------------------------------
    head = PatternFill("solid", fgColor="1F3A5F")

    def table(title, headers, rows, widths):
        s = wb.create_sheet(title)
        s.append(headers)
        for c in s[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), head
        for row in rows:
            s.append(row)
        for n, w in enumerate(widths, 1):
            s.column_dimensions[openpyxl.utils.get_column_letter(n)].width = w
        for row in s.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        s.freeze_panes = "A2"
        _fit_width(s, landscape=True)
        return s

    table("AI Review", ["Item", "Value"], [
        ["Status", "AI DRAFT. Not a title opinion. A landman must check every item and sign off."],
        ["Run", str(run_id)], ["Tract", intake.tract_id],
        ["Package source", intake.mode + (f" ({intake.runsheet_path.name})" if intake.runsheet_path else "")],
        ["Tract acreage", f"{acres:g}" if acres else "unknown (add to tracts.csv)"],
        ["Current surface owner", names(owners) or "NOT DETERMINED"],
        ["Vesting instrument", display_ref(vesting) if vesting else None],
        ["Flags", str(len(flags))],
        ["Ownership notes", "\n".join(m for _, m in chain.get("ownership_notes", [])) or "none"],
        ["Instruments with an unrecognized type", ", ".join(chain["unclassified"]) or "none"],
        ["Instruments with no filing date", ", ".join(chain["undated"]) or "none"],
        ["Rules not yet implemented", ", ".join(skipped) or "none"],
        ["Package problems", "\n".join(problems) or "none"],
        ["Red dates", "On the Index and Runsheet, a red file date is earlier than the effective date: check it."],
    ], [34, 100])
    fs = table("AI Flags", ["Rule", "Severity", "Instrument", "Message", "Reviewer action", "Note"],
               [[f.rule_id, f.severity, f.instrument_id, f.message, None, None] for f in flags], [8, 12, 26, 80, 16, 40])
    colors = {"high": "F8D7DA", "medium": "FFF2CC", "low": "E2EFDA", "review": "DDEBF7"}
    for row in fs.iter_rows(min_row=2):
        color = colors.get(str(row[1].value or "").split()[0].lower() if row[1].value else "")
        if color:
            row[1].fill = PatternFill("solid", fgColor=color)
    table("AI Supporting Docs", ["Kind", "File"], [[i.get("kind"), i["file_name"]] for i in chain.get("supporting", [])],
          [30, 80])

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.active = 0
    wb.save(path)
    return path
